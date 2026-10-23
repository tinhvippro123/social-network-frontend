import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report_docx():
    doc = docx.Document()

    # Standard Vietnamese Thesis Margins: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # ~2.0 cm
        section.bottom_margin = Inches(0.79)  # ~2.0 cm
        section.left_margin = Inches(1.18)    # ~3.0 cm
        section.right_margin = Inches(0.79)   # ~2.0 cm

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(13)
    normal_style.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    normal_style.paragraph_format.line_spacing = 1.3
    normal_style.paragraph_format.space_after = Pt(4)

    md_path = os.path.join(os.path.dirname(__file__), 'BAO_CAO_DO_AN_CHUYEN_NGANH.md')
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found")
        return

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle page breaks
        if '\\newpage' in stripped:
            doc.add_page_break()
            continue

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            # End of table, process it
            in_table = False
            process_table(doc, table_lines)
            table_lines = []

        if not stripped:
            continue

        # Skip center div tags
        if stripped in ['<div align="center">', '</div>', '<br>', '<br><br>']:
            continue

        # Heading 1
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if 'TRƯỜNG ĐẠI HỌC' in text or 'BÁO CÁO' in text or 'XÂY DỰNG' in text:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Navy blue
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)

        # Heading 2
        elif stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)

        # Heading 3
        elif stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)

        # Blockquote
        elif stripped.startswith('> '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.2)
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

        # Bullet list
        elif stripped.startswith('- ') or stripped.startswith('* '):
            text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_text(p, text)

        # Numbered list
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped).strip()
            p = doc.add_paragraph(style='List Number')
            add_formatted_text(p, text)

        # Normal paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.3)
            add_formatted_text(p, stripped)

    if in_table and table_lines:
        process_table(doc, table_lines)

    output_path = os.path.join(os.path.dirname(__file__), '1250080199_LeThanhTinh_12_DH_CNPM3_BaoCaoDACN.docx')
    doc.save(output_path)
    print(f"SUCCESS: Report saved to {output_path}")

def add_formatted_text(paragraph, text):
    # Regex to split bold (**text**) and plain text
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.font.name = 'Times New Roman'
            run.font.bold = True
        else:
            # Handle italics (*text*)
            sub_parts = re.split(r'(\*.*?\*)', part)
            for sub in sub_parts:
                if sub.startswith('*') and sub.endswith('*') and len(sub) > 2:
                    run = paragraph.add_run(sub[1:-1])
                    run.font.name = 'Times New Roman'
                    run.font.italic = True
                else:
                    run = paragraph.add_run(sub)
                    run.font.name = 'Times New Roman'

def process_table(doc, lines):
    # Filter separator lines like |---|:---:|
    valid_rows = [l for l in lines if not re.match(r'^\|[\s\-:]+\|\s*$', l)]
    if not valid_rows:
        return

    data = []
    for r in valid_rows:
        cells = [c.strip() for c in r.strip('|').split('|')]
        data.append(cells)

    if not data:
        return

    cols = len(data[0])
    table = doc.add_table(rows=len(data), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, row in enumerate(data):
        is_header = (i == 0)
        for j, cell_text in enumerate(row):
            if j < cols:
                cell = table.cell(i, j)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                
                if is_header:
                    set_cell_background(cell, '1E3A8A') # Navy blue
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(cell_text.replace('**', ''))
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if i % 2 == 1:
                        set_cell_background(cell, 'F8FAFC') # subtle light grey
                    run = p.add_run(cell_text.replace('**', ''))
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    if ':---:' in lines[1] if len(lines) > 1 else False:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph() # Add spacing after table

if __name__ == '__main__':
    create_report_docx()
